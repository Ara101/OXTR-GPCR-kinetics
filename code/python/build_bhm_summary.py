"""
Write the BHM refit (code/python/bhm_refit.py) into data/GPCR_parameters.xlsx and compute Tables S2/S3.

data/GPCR_parameters.xlsx was maintained by hand. This script replaces only its BHM columns (REM
columns are untouched) with the refit's posterior medians and 95% HDIs from
results/bhm_refit_2026-09-24/tables/bhm_refit_summary.csv, rounded to three significant figures as
reported in Table S1. Before overwriting, it records the values being replaced (the 2026-09-23 rerun
of the original notebook) next to the new ones in bhm_refit_vs_previous.csv.

Table S2/S3 statistics use the formula of 1_comparative_analysis_bhm_rem.ipynb (cell 4):
SE = (ln U - ln L) / (2 * 1.959964) for each model, Z = (ln m_BHM - ln m_REM) / sqrt(SE_BHM^2 + SE_REM^2),
two-sided p, effect size = Z * sqrt(2). Inputs are the rounded values printed in the tables.
"""

from math import erf, log, sqrt
from pathlib import Path

import openpyxl
import pandas as pd

REPO = Path(__file__).resolve().parents[2]
RUN = REPO / "results" / "bhm_refit_2026-09-24"
BOOK = REPO / "data" / "GPCR_parameters.xlsx"
BHM_SHEETS = ["GPCR_parameters_geomean", "GPCR_parameters_beta12_geomean"]
# summary-sheet row name -> refit summary Parameter
ROWS = {"Kd-OXTR: G-protein": "Kd", "konG": "konG", "koffG": "koffG", "kact": "kact", "kinact": "kinact", "kP": "kP",
        "kdp": "kdp", "kba": "kba", "kba 1": "kba1", "kba 2": "kba2", "kint": "kint",
        "kcons_int": "kcons_int", "krec": "krec", "kexpt": "kexpt", "krgt": "krgt", "kdeg": "kdeg"}
COLS = ["BHM (mean)", "BHM (lower [95% CI])", "BHM (upper [95% CI])"]


def sig3(x):
    return float(f"{x:.3g}")


def main():
    refit = pd.read_csv(RUN / "tables" / "bhm_refit_summary.csv").set_index("Parameter")
    values = {row: tuple(sig3(refit.loc[p, c]) for c in ("median", "hdi_lo", "hdi_hi")) for row, p in ROWS.items()}

    # record what is being replaced (geomean sheets hold the values the documents report)
    geo_old = pd.read_excel(BOOK, sheet_name="GPCR_parameters_geomean").set_index("Parameter")
    b12_old = pd.read_excel(BOOK, sheet_name="GPCR_parameters_beta12_geomean").set_index("Parameter")
    comp = []
    for row, p in ROWS.items():
        src = b12_old if row in ("kba 1", "kba 2") else geo_old
        if row not in src.index:
            continue
        old = [sig3(float(src.loc[row, c])) for c in COLS]
        new = values[row]
        comp.append(dict(Parameter=p, old_mean=old[0], old_lo=old[1], old_hi=old[2], new_mean=new[0], new_lo=new[1],
                         new_hi=new[2], fold_new_over_old=new[0] / old[0], status="ok"))
    out = RUN / "bhm_refit_vs_previous.csv"
    if not out.exists():  # keep the first comparison (previous = 2026-09-23 rerun) if the script is run again
        pd.DataFrame(comp).to_csv(out, index=False)

    wb = openpyxl.load_workbook(BOOK)
    for sheet in BHM_SHEETS:
        for r in wb[sheet].iter_rows(min_row=2):
            if r[0].value in values:
                for cell, v in zip(r[1:4], values[r[0].value]):
                    cell.value = v
    ws = wb["beta_arrestin_bhm"]  # notebook-format sheet: median and HDI only; SD/quartiles are not produced by the refit
    for r in ws.iter_rows(min_row=2):
        key = {"beta_arrestin_intercept[0]": "kba 1", "beta_arrestin_intercept[1]": "kba 2"}.get(r[0].value)
        if key:
            r[1].value, r[3].value, r[4].value = values[key]
            r[2].value = r[5].value = r[6].value = None
        elif r[0].value == "Global Mean":
            for c in r[1:7]:
                c.value = None
    wb.save(BOOK)

    geo = pd.read_excel(BOOK, sheet_name="GPCR_parameters_geomean").set_index("Parameter")
    b12 = pd.read_excel(BOOK, sheet_name="GPCR_parameters_beta12_geomean").set_index("Parameter")
    rows = []
    for p, src in [("kba 1", b12), ("kba 2", b12), ("kba", geo), ("kact", geo), ("kinact", geo), ("kP", geo), ("kdp", geo),
                   ("kcons_int", geo), ("kint", geo), ("krec", geo), ("kexpt", geo), ("krgt", geo), ("kdeg", geo)]:
        b = [float(src.loc[p, c]) for c in COLS]
        r = [float(src.loc[p, c]) for c in ["REM (mean)", "REM (lower [95% CI])", "REM (upper [95% CI])"]]
        b, r = [sig3(x) for x in b], [sig3(x) for x in r]
        se_b = (log(b[2]) - log(b[1])) / (2 * 1.959964)
        se_r = (log(r[2]) - log(r[1])) / (2 * 1.959964)
        z = (log(b[0]) - log(r[0])) / sqrt(se_b ** 2 + se_r ** 2)
        pval = 2 * (1 - 0.5 * (1 + erf(abs(z) / sqrt(2))))
        rows.append(dict(Parameter=p, BHM=b[0], BHM_lo=b[1], BHM_hi=b[2], REM=r[0], REM_lo=r[1], REM_hi=r[2],
                         SE_BHM_log=se_b, SE_REM_log=se_r, Z=z, p=pval, significant=pval < 0.05, effect_size=z * sqrt(2)))
    stats = pd.DataFrame(rows)
    stats.to_csv(RUN / "table_s2_s3_statistics.csv", index=False)
    with pd.option_context("display.float_format", "{:.3g}".format, "display.width", 220):
        print(pd.DataFrame(comp).to_string(index=False))
        print(stats.to_string(index=False))


if __name__ == "__main__":
    main()
